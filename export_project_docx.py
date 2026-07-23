"""Ekspor PROJECT_DOCUMENTATION.md menjadi dokumen Word yang rapi."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "PROJECT_DOCUMENTATION.md"
OUTPUT = ROOT / "Panduan_AI_Trading_Bot_untuk_Pemula_v1.2.2.docx"


def set_cell_shading(cell, color: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    tc_pr.append(shading)


def add_page_number(paragraph) -> None:
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Halaman ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    run._r.addnext(field)


def add_inline(paragraph, text: str) -> None:
    """Tambahkan markdown inline sederhana: bold dan inline code."""
    parts = re.split(r"(`[^`]+`|\*\*[^*]+\*\*)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(30, 64, 175)
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        else:
            paragraph.add_run(part)


def configure_document(document: Document) -> None:
    section = document.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.0)

    normal = document.styles["Normal"]
    normal.font.name = "Aptos"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.12

    for name, size, color in (
        ("Title", 24, RGBColor(15, 23, 42)),
        ("Heading 1", 17, RGBColor(30, 64, 175)),
        ("Heading 2", 13, RGBColor(15, 118, 110)),
        ("Heading 3", 11, RGBColor(51, 65, 85)),
    ):
        style = document.styles[name]
        style.font.name = "Aptos Display"
        style.font.size = Pt(size)
        style.font.color.rgb = color

    if "Code Block" not in document.styles:
        code = document.styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
        code.font.name = "Consolas"
        code.font.size = Pt(8.5)
        code.font.color.rgb = RGBColor(226, 232, 240)
        code.paragraph_format.left_indent = Cm(0.4)
        code.paragraph_format.right_indent = Cm(0.4)
        code.paragraph_format.space_before = Pt(4)
        code.paragraph_format.space_after = Pt(4)

    footer = section.footer.paragraphs[0]
    add_page_number(footer)


def add_code_block(document: Document, lines: list[str]) -> None:
    paragraph = document.add_paragraph(style="Code Block")
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "0F172A")
    paragraph._p.get_or_add_pPr().append(shading)
    run = paragraph.add_run("\n".join(lines))
    run.font.name = "Consolas"


def add_table(document: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    table = document.add_table(rows=len(rows), cols=column_count)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for row_index, row in enumerate(rows):
        for column_index in range(column_count):
            cell = table.cell(row_index, column_index)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            text = row[column_index] if column_index < len(row) else ""
            paragraph = cell.paragraphs[0]
            add_inline(paragraph, text)
            if row_index == 0:
                set_cell_shading(cell, "1E40AF")
                for run in paragraph.runs:
                    run.bold = True
                    run.font.color.rgb = RGBColor(255, 255, 255)
            elif row_index % 2 == 0:
                set_cell_shading(cell, "E2E8F0")
    document.add_paragraph()


def export() -> Path:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    document = Document()
    configure_document(document)

    document.add_heading("Panduan Lengkap\nAI Trading Bot untuk Pemula", 0)
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("Versi Desktop 1.2.2").bold = True
    subtitle.add_run("\nPenjelasan sederhana dari dasar sampai seluruh pengaturan")
    document.add_paragraph()
    notice = document.add_paragraph()
    notice.alignment = WD_ALIGN_PARAGRAPH.CENTER
    notice.add_run("Dokumen proyek • 23 Juli 2026").italic = True
    document.add_page_break()

    # Lewati judul Markdown karena sudah dibuat sebagai cover.
    index = 1 if lines and lines[0].startswith("# ") else 0
    in_code = False
    code_lines: list[str] = []
    while index < len(lines):
        line = lines[index]
        stripped = line.strip()

        if stripped.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue
        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            table_rows: list[list[str]] = []
            while index < len(lines):
                candidate = lines[index].strip()
                if not (candidate.startswith("|") and candidate.endswith("|")):
                    break
                cells = [cell.strip() for cell in candidate.strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells):
                    table_rows.append(cells)
                index += 1
            add_table(document, table_rows)
            continue

        heading = re.match(r"^(#{2,4})\s+(.+)$", stripped)
        if heading:
            level = min(len(heading.group(1)) - 1, 3)
            document.add_heading(heading.group(2), level=level)
            index += 1
            continue

        if stripped.startswith("> "):
            paragraph = document.add_paragraph(style="Intense Quote")
            add_inline(paragraph, stripped[2:])
        elif re.match(r"^[-*]\s+", stripped):
            paragraph = document.add_paragraph(style="List Bullet")
            add_inline(paragraph, re.sub(r"^[-*]\s+", "", stripped))
        elif re.match(r"^\d+\.\s+", stripped):
            paragraph = document.add_paragraph(style="List Number")
            add_inline(paragraph, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped:
            paragraph = document.add_paragraph()
            add_inline(paragraph, stripped)

        index += 1

    if in_code and code_lines:
        add_code_block(document, code_lines)

    # Metadata inti dokumen.
    properties = document.core_properties
    properties.title = "Panduan Lengkap AI Trading Bot untuk Pemula"
    properties.subject = "AI Trading Desktop v1.2.2"
    properties.author = "AI Trading Project"
    properties.keywords = "MetaTrader 5, AI trading, risk management, machine learning"

    document.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(export())
